import re
import streamlit as st
import io
import base64
import json
import os
import random
from datetime import datetime

try:
    import PyPDF2
except ImportError:
    st.error("PyPDF2 is not installed. Please install it with: pip install PyPDF2")

try:
    import docx
except ImportError:
    st.error("python-docx is not installed. Please install it with: pip install python-docx")

from fpdf import FPDF
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication

try:
    import nltk
    from nltk.tokenize import word_tokenize
    from nltk.corpus import stopwords
    from nltk.stem import WordNetLemmatizer
    
    nltk.download('punkt', quiet=True)
    nltk.download('stopwords', quiet=True)
    nltk.download('wordnet', quiet=True)
    
    NLP_ENABLED = True
except ImportError as e:
    st.warning(f"NLTK not installed: {e}. NLP features will be disabled.")
    NLP_ENABLED = False

COMMON_SKILLS = {
    'programming': ['python', 'java', 'javascript', 'html', 'css', 'c++', 'c#', 'ruby', 'php', 'sql', 'r'],
    'frameworks': ['react', 'angular', 'vue', 'django', 'flask', 'spring', 'node.js', 'express', '.net'],
    'databases': ['sql', 'mysql', 'postgresql', 'mongodb', 'oracle', 'sqlite', 'redis'],
    'cloud': ['aws', 'azure', 'gcp', 'google cloud', 'docker', 'kubernetes'],
    'tools': ['git', 'github', 'jira', 'jenkins', 'agile', 'scrum'],
    'soft_skills': ['communication', 'leadership', 'teamwork', 'problem solving', 'time management'],
}

# Expanded technical questions for each skill
TECHNICAL_QUESTIONS = {
    'python': [
        {"question": "Explain how you would implement a decorator in Python.", 
         "expected_keywords": ["function", "wrapper", "decorator", "@", "arguments", "return"]},
        {"question": "How would you handle exceptions in Python?", 
         "expected_keywords": ["try", "except", "finally", "raise", "error", "handling"]},
        {"question": "Describe the difference between a list and a tuple in Python.", 
         "expected_keywords": ["mutable", "immutable", "list", "tuple", "ordered", "elements"]},
        {"question": "How do you handle concurrent operations in Python?",
         "expected_keywords": ["threading", "multiprocessing", "GIL", "asyncio", "concurrent", "futures"]},
        {"question": "Explain Python's memory management and garbage collection.",
         "expected_keywords": ["reference counting", "garbage collector", "memory", "objects", "allocation", "cycle"]},
        {"question": "What are Python generators and how do they work?",
         "expected_keywords": ["yield", "iterator", "memory", "efficient", "lazy", "next"]},
        {"question": "How would you optimize a slow Python script?",
         "expected_keywords": ["profiling", "algorithms", "data structures", "cython", "numpy", "vectorization"]},
        {"question": "Explain the difference between shallow copy and deep copy in Python.",
         "expected_keywords": ["copy", "deepcopy", "reference", "nested", "objects", "mutable"]},
        {"question": "How do you use context managers in Python?",
         "expected_keywords": ["with", "context", "enter", "exit", "resources", "cleanup"]},
        {"question": "Describe Python's approach to OOP principles.",
         "expected_keywords": ["class", "inheritance", "polymorphism", "encapsulation", "self", "method"]}
    ],
    'java': [
        {"question": "Explain the concept of inheritance in Java.", 
         "expected_keywords": ["extends", "class", "parent", "child", "super", "override"]},
        {"question": "How do you handle exceptions in Java?", 
         "expected_keywords": ["try", "catch", "finally", "throw", "throws", "exception"]},
        {"question": "What is the difference between an interface and an abstract class in Java?", 
         "expected_keywords": ["implement", "extend", "methods", "abstract", "interface", "multiple"]},
        {"question": "Explain Java's memory management and garbage collection.",
         "expected_keywords": ["heap", "stack", "JVM", "garbage collection", "memory", "objects"]},
        {"question": "How do you handle concurrency in Java?",
         "expected_keywords": ["synchronized", "thread", "lock", "atomic", "concurrent", "volatile"]},
        {"question": "What are Java streams and how do you use them?",
         "expected_keywords": ["stream", "filter", "map", "collect", "functional", "lambda"]},
        {"question": "Explain the principle of 'Write Once, Run Anywhere' in Java.",
         "expected_keywords": ["JVM", "bytecode", "platform", "independent", "compile", "interpret"]},
        {"question": "How do annotations work in Java?",
         "expected_keywords": ["annotation", "metadata", "runtime", "compiler", "reflection", "processor"]},
        {"question": "What are the different types of references in Java?",
         "expected_keywords": ["strong", "weak", "soft", "phantom", "garbage", "collection"]},
        {"question": "Explain Java generics and their benefits.",
         "expected_keywords": ["generics", "type", "safety", "compile", "erasure", "collections"]}
    ],
    'javascript': [
        {"question": "Explain closures in JavaScript.", 
         "expected_keywords": ["function", "scope", "variable", "closure", "lexical", "access"]},
        {"question": "How does asynchronous programming work in JavaScript?", 
         "expected_keywords": ["promise", "async", "await", "callback", "then", "event loop"]},
        {"question": "What's the difference between var, let, and const in JavaScript?", 
         "expected_keywords": ["scope", "hoisting", "reassign", "block", "function", "declaration"]},
        {"question": "Explain prototypal inheritance in JavaScript.",
         "expected_keywords": ["prototype", "inheritance", "chain", "object", "__proto__", "constructor"]},
        {"question": "How does the 'this' keyword work in JavaScript?",
         "expected_keywords": ["context", "bind", "call", "apply", "arrow", "function"]},
        {"question": "What are JavaScript modules and how do you use them?",
         "expected_keywords": ["import", "export", "module", "ESM", "CommonJS", "bundler"]},
        {"question": "Explain event bubbling and capturing in JavaScript.",
         "expected_keywords": ["bubbling", "capturing", "propagation", "target", "currentTarget", "stopPropagation"]},
        {"question": "How does JavaScript's event loop work?",
         "expected_keywords": ["stack", "queue", "event loop", "microtask", "macrotask", "asynchronous"]},
        {"question": "What are Web APIs and how do they interact with JavaScript?",
         "expected_keywords": ["DOM", "fetch", "localStorage", "browser", "API", "asynchronous"]},
        {"question": "Describe JavaScript's memory management and garbage collection.",
         "expected_keywords": ["garbage collection", "memory leak", "reference", "closure", "heap", "mark-and-sweep"]}
    ],
    'sql': [
        {"question": "Explain the difference between INNER JOIN and LEFT JOIN.", 
         "expected_keywords": ["inner", "left", "join", "matching", "all", "records"]},
        {"question": "How would you optimize a slow SQL query?", 
         "expected_keywords": ["index", "execution plan", "query", "optimize", "performance", "analyze"]},
        {"question": "What is database normalization?", 
         "expected_keywords": ["normal form", "redundancy", "dependency", "relation", "table", "normalize"]},
        {"question": "Explain the concept of transactions in SQL.",
         "expected_keywords": ["ACID", "commit", "rollback", "transaction", "atomic", "consistency"]},
        {"question": "How do you handle deadlocks in a database?",
         "expected_keywords": ["deadlock", "transaction", "lock", "timeout", "detection", "prevention"]},
        {"question": "What are indexes and how do they improve query performance?",
         "expected_keywords": ["index", "b-tree", "search", "performance", "clustered", "non-clustered"]},
        {"question": "Explain the difference between DELETE, TRUNCATE, and DROP.",
         "expected_keywords": ["delete", "truncate", "drop", "table", "record", "structure"]},
        {"question": "How do you implement data partitioning in SQL?",
         "expected_keywords": ["partition", "shard", "horizontal", "vertical", "table", "performance"]},
        {"question": "What are stored procedures and triggers?",
         "expected_keywords": ["stored procedure", "trigger", "function", "execute", "event", "automate"]},
        {"question": "Explain subqueries and their types in SQL.",
         "expected_keywords": ["subquery", "nested", "correlated", "uncorrelated", "performance", "optimization"]}
    ],
    'react': [
        {"question": "Explain the component lifecycle in React.", 
         "expected_keywords": ["mount", "update", "unmount", "render", "effect", "component"]},
        {"question": "How do you manage state in React applications?", 
         "expected_keywords": ["useState", "useReducer", "state", "props", "context", "Redux"]},
        {"question": "What are hooks in React and why were they introduced?", 
         "expected_keywords": ["hooks", "functional", "state", "effect", "rules", "useState"]},
        {"question": "Explain the concept of Virtual DOM in React.",
         "expected_keywords": ["virtual DOM", "reconciliation", "diffing", "rendering", "performance", "update"]},
        {"question": "How do you handle side effects in React components?",
         "expected_keywords": ["useEffect", "dependency array", "cleanup", "lifecycle", "async", "fetch"]},
        {"question": "What are Higher-Order Components (HOCs) in React?",
         "expected_keywords": ["HOC", "component", "wrapping", "reuse", "props", "enhance"]},
        {"question": "Explain the Context API and when you would use it.",
         "expected_keywords": ["context", "provider", "consumer", "useContext", "global", "state"]},
        {"question": "How do you optimize performance in React applications?",
         "expected_keywords": ["memo", "useMemo", "useCallback", "lazy loading", "code splitting", "performance"]},
        {"question": "What are React portals and when would you use them?",
         "expected_keywords": ["portal", "DOM", "render", "modal", "overlay", "parent"]},
        {"question": "Explain the difference between controlled and uncontrolled components.",
         "expected_keywords": ["controlled", "uncontrolled", "input", "form", "state", "ref"]}
    ],
    'aws': [
        {"question": "Explain the difference between EC2 and Lambda.", 
         "expected_keywords": ["instance", "serverless", "EC2", "Lambda", "scaling", "compute"]},
        {"question": "How do you handle security in AWS?", 
         "expected_keywords": ["IAM", "security group", "encryption", "access", "policy", "role"]},
        {"question": "Describe the AWS services you've worked with.", 
         "expected_keywords": ["S3", "EC2", "Lambda", "RDS", "CloudFront", "DynamoDB"]},
        {"question": "What is Auto Scaling in AWS and how does it work?",
         "expected_keywords": ["scaling", "EC2", "group", "load", "policy", "metric"]},
        {"question": "Explain AWS VPC and its components.",
         "expected_keywords": ["VPC", "subnet", "route table", "internet gateway", "NACL", "security group"]},
        {"question": "How do you deploy applications in AWS?",
         "expected_keywords": ["CodeDeploy", "Elastic Beanstalk", "CloudFormation", "pipeline", "CI/CD", "deployment"]},
        {"question": "What are AWS S3 storage classes and when would you use each?",
         "expected_keywords": ["standard", "infrequent access", "glacier", "cost", "retrieval", "durability"]},
        {"question": "Explain AWS CloudFormation and Infrastructure as Code.",
         "expected_keywords": ["template", "stack", "resource", "IaC", "provision", "automation"]},
        {"question": "How do you monitor applications in AWS?",
         "expected_keywords": ["CloudWatch", "logging", "metrics", "alarm", "event", "dashboard"]},
        {"question": "What is AWS Identity and Access Management (IAM)?",
         "expected_keywords": ["IAM", "user", "role", "policy", "permission", "authentication"]}
    ],
}

# Expanded generic questions
GENERIC_QUESTIONS = [
    {"question": "Tell me about a challenging project you worked on and how you overcame obstacles.", 
     "expected_keywords": ["challenge", "project", "solution", "overcome", "team", "result"]},
    {"question": "How do you approach learning new technologies?", 
     "expected_keywords": ["learning", "research", "practice", "curiosity", "documentation", "projects"]},
    {"question": "Describe your experience with agile development methodologies.", 
     "expected_keywords": ["agile", "scrum", "sprint", "kanban", "standup", "retrospective"]},
    {"question": "How do you ensure code quality in your projects?", 
     "expected_keywords": ["testing", "review", "standards", "documentation", "refactoring", "clean"]},
    {"question": "Describe a situation where you had to debug a complex issue. What was your approach?",
     "expected_keywords": ["debug", "troubleshoot", "analyze", "problem", "solution", "methodology"]},
    {"question": "How do you keep your technical skills up-to-date?",
     "expected_keywords": ["learning", "courses", "practice", "community", "projects", "research"]},
    {"question": "Tell me about a time when you had to make a technical decision with limited information.",
     "expected_keywords": ["decision", "analysis", "risk", "information", "outcome", "process"]},
    {"question": "How do you handle technical disagreements within a team?",
     "expected_keywords": ["communication", "compromise", "discussion", "evidence", "respect", "resolution"]},
    {"question": "Describe your approach to system design and architecture.",
     "expected_keywords": ["scalability", "requirements", "tradeoffs", "design", "components", "architecture"]},
    {"question": "How do you balance technical debt against delivering features?",
     "expected_keywords": ["technical debt", "prioritize", "refactor", "balance", "quality", "delivery"]},
    {"question": "Tell me about a time when you improved a process or system.",
     "expected_keywords": ["improvement", "efficiency", "process", "impact", "measure", "implement"]},
    {"question": "How do you approach documentation in your projects?",
     "expected_keywords": ["documentation", "clarity", "audience", "purpose", "update", "importance"]}
]

WELCOME_MESSAGES = [
    "Welcome to TechInterviewBot! I'm here to help you practice your technical interview skills.",
    "Hello! I'm your Technical Interview Assistant. Let's prepare you for your next tech interview.",
    "Hi there! Ready to sharpen your technical interview skills? I'm here to help you practice.",
    "Welcome aboard! I'm your AI interview coach. Let's see how well you can handle technical questions."
]

RESUME_PROMPTS = [
    "To get started, please upload your resume or paste its content so I can tailor questions to your skills.",
    "Let's begin by analyzing your resume. Please upload it or paste the text so I can customize the interview.",
    "First, I'll need to see your resume to generate relevant questions. Upload a file or paste the text below.",
    "To create a personalized interview experience, I need to review your resume first. Upload or paste it below."
]

SKILL_MESSAGES = [
    "Great! I've analyzed your resume and identified these key skills:",
    "Thanks for sharing your resume! Based on my analysis, here are the skills I've identified:",
    "Perfect! After reviewing your resume, I've extracted these technical skills:",
    "I've processed your resume and found these skills that we can focus on:"
]

INTERVIEW_START_MESSAGES = [
    "Now let's begin the interview! I'll ask you a series of technical questions related to your skills.",
    "Ready to start? I've prepared some technical questions based on your experience.",
    "Let's dive into the technical interview! I'll ask questions related to your strongest skills.",
    "The interview is about to begin! I'll evaluate your answers to help improve your skills."
]

QUESTION_TRANSITIONS = [
    "Let's move on to the next question:",
    "Here's another question for you:",
    "Now, I'd like to ask you about:",
    "Let's continue with this question:",
    "For the next question:",
    "Moving forward:",
]

EVALUATION_POSITIVE = [
    "Great answer! You've covered the key points effectively.",
    "Excellent response! Your understanding of the concept is clear.",
    "Well done! Your explanation was thorough and accurate.",
    "Very good! You demonstrated strong knowledge in this area."
]

EVALUATION_AVERAGE = [
    "Good attempt! You covered some key points, but there's room for improvement.",
    "That's a decent answer, but you could expand on a few concepts.",
    "Not bad! You have the basic understanding, but consider adding more depth.",
    "You're on the right track, but try to be more specific in your explanations."
]

EVALUATION_NEEDS_IMPROVEMENT = [
    "You've made an attempt, but there are some key concepts missing.",
    "Your answer needs more technical depth. Let me suggest some areas to focus on.",
    "I see you have some understanding, but there are important points you didn't address.",
    "This response could be improved by including more specific technical details."
]

def preprocess_text(text):
    if not NLP_ENABLED:
        return text.lower()
    tokens = word_tokenize(text.lower())
    stop_words = set(stopwords.words('english'))
    tokens = [token for token in tokens if token not in stop_words and token.isalnum()]
    lemmatizer = WordNetLemmatizer()
    tokens = [lemmatizer.lemmatize(token) for token in tokens]
    return " ".join(tokens)

def extract_skills_with_nlp(text):
    """Extract skills with very strict context-based matching."""
    if not NLP_ENABLED:
        return extract_skills_basic(text)
    
    raw_text = text.lower()
    identified_skills = {}
    debug_matches = []
    
    SKILL_ALIASES = {
        'python': ['python'],
        'java': ['java'],
        'javascript': ['javascript'],
        'sql': ['sql'],
        'aws': ['aws'],
        'react': ['react'],
        'django': ['django'],
        'flask': ['flask'],
        'git': ['git'],
        'github': ['github'],
    }
    
    context_patterns = r'(?:skills|experience|proficient in|worked with|knowledge of|using|expertise in|developed with)'
    
    for category, skill_list in COMMON_SKILLS.items():
        found_skills = set()
        for skill in skill_list:
            aliases = SKILL_ALIASES.get(skill, [skill])
            for alias in aliases:
                # Require exact match near context keyword
                pattern = rf'{context_patterns}\s*[^.\n]*\b{re.escape(alias)}\b[^.\n]*'
                matches = re.finditer(pattern, raw_text)
                for match in matches:
                    found_skills.add(skill)
                    debug_matches.append(f"Matched '{skill}' in: '{match.group()}'")
        if found_skills:
            identified_skills[category] = list(found_skills)
    
    st.session_state.debug_skills = debug_matches
    st.session_state.raw_resume_text = raw_text  # Store raw text for debugging
    return identified_skills

def extract_skills_basic(text):
    """Basic skill extraction with strict context."""
    text = text.lower()
    identified_skills = {}
    context_patterns = r'(?:skills|experience|proficient in|worked with|knowledge of|using|expertise in|developed with)'
    
    for category, skill_list in COMMON_SKILLS.items():
        found_skills = set()
        for skill in skill_list:
            pattern = rf'{context_patterns}\s*[^.\n]*\b{re.escape(skill)}\b[^.\n]*'
            if re.search(pattern, text):
                found_skills.add(skill)
        if found_skills:
            identified_skills[category] = list(found_skills)
    return identified_skills

def extract_skills(text):
    if not text:
        return {}
    return extract_skills_with_nlp(text) if NLP_ENABLED else extract_skills_basic(text)

def evaluate_answer_with_nlp(question, answer, expected_keywords):
    """Evaluate an answer using NLP techniques"""
    if not answer.strip():
        return {"score": 0, "feedback": "No answer provided.", "missing_concepts": expected_keywords}
    
    # Basic keyword matching if NLP is not enabled
    if not NLP_ENABLED:
        keyword_count = sum(1 for keyword in expected_keywords if keyword.lower() in answer.lower())
        score = min(keyword_count / len(expected_keywords), 1.0) * 100
        missing = [k for k in expected_keywords if k.lower() not in answer.lower()]
        return {"score": score, "feedback": "Basic keyword matching applied.", "missing_concepts": missing}
    
    # Advanced NLP processing
    processed_answer = preprocess_text(answer)
    processed_keywords = [preprocess_text(kw) for kw in expected_keywords]
    
    # Count matched keywords
    keyword_count = sum(1 for kw in processed_keywords if kw in processed_answer)
    
    # Calculate score based on keyword coverage
    score = min(keyword_count / len(expected_keywords), 1.0) * 100
    
    # Identify missing concepts
    missing = [kw for kw in expected_keywords if preprocess_text(kw) not in processed_answer]
    
    # Generate feedback based on score
    feedback = get_feedback_message(score)
    if missing:
        feedback += f" Consider mentioning: {', '.join(missing[:3])}."
    
    # For technical depth analysis, check sentence length and variation
    sentences = answer.split('.')
    avg_sentence_length = sum(len(s.split()) for s in sentences if s.strip()) / max(1, len([s for s in sentences if s.strip()]))
    
    # Analyze technical depth based on answer length and keyword density
    if len(answer.split()) > 100 and score > 70:
        feedback += " Your answer shows good technical depth."
    elif len(answer.split()) < 30 and score < 70:
        feedback += " Consider providing a more detailed explanation."
    
    # Check for explanation patterns
    explanation_patterns = ["because", "therefore", "means that", "this is why", "which is"]
    has_explanations = any(pattern in answer.lower() for pattern in explanation_patterns)
    
    if has_explanations and score > 50:
        feedback += " Your explanations help demonstrate understanding."
    elif not has_explanations and score > 50:
        feedback += " Consider explaining 'why' in addition to 'what'."
    
    return {
        "score": score, 
        "feedback": feedback, 
        "missing_concepts": missing
    }

def extract_text_from_pdf(pdf_file):
    try:
        pdf_reader = PyPDF2.PdfReader(pdf_file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text() or ""
        return text
    except Exception as e:
        st.error(f"Error processing PDF: {e}")
        return ""

def extract_text_from_docx(docx_file):
    try:
        doc = docx.Document(docx_file)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text
    except Exception as e:
        st.error(f"Error processing DOCX: {e}")
        return ""

def generate_technical_questions(skills, max_questions=7):
    """Generate questions with improved randomization and more questions per skill"""
    all_possible_questions = []
    
    # Get all skills from the user's profile
    all_skills = [skill for category, skill_list in skills.items() for skill in skill_list]
    
    # Count frequency of each skill
    skill_frequency = {skill: all_skills.count(skill) for skill in set(all_skills)}
    
    # Sort skills by frequency (most frequent first)
    sorted_skills = sorted(skill_frequency.keys(), key=lambda x: skill_frequency[x], reverse=True)
    
    # Aim to select a balanced set of questions
    target_questions_per_skill = max(1, max_questions // len(sorted_skills)) if sorted_skills else 0
    remaining_slots = max_questions
    
    # Collect questions for each skill
    for skill in sorted_skills:
        if skill in TECHNICAL_QUESTIONS:
            # Get all questions for this skill and shuffle them
            skill_questions = TECHNICAL_QUESTIONS[skill].copy()
            random.shuffle(skill_questions)
            
            # Take up to target number of questions per skill
            num_to_take = min(len(skill_questions), target_questions_per_skill, remaining_slots)
            all_possible_questions.extend(skill_questions[:num_to_take])
            remaining_slots -= num_to_take
    
    # If we still have slots available, add random questions from any skill
    if remaining_slots > 0 and sorted_skills:
        additional_questions = []
        for skill in sorted_skills:
            if skill in TECHNICAL_QUESTIONS:
                # Get questions we haven't used yet
                used_questions = [q["question"] for q in all_possible_questions]
                unused_questions = [q for q in TECHNICAL_QUESTIONS[skill] 
                                  if q["question"] not in used_questions]
                additional_questions.extend(unused_questions)
        
        # Shuffle and take what we need
        random.shuffle(additional_questions)
        all_possible_questions.extend(additional_questions[:remaining_slots])
        remaining_slots -= min(len(additional_questions), remaining_slots)
    
    # Add generic questions if we still need more
    if remaining_slots > 0:
        random.shuffle(GENERIC_QUESTIONS)
        
        # Get questions we haven't used yet
        used_questions = [q["question"] for q in all_possible_questions]
        unused_generic = [q for q in GENERIC_QUESTIONS if q["question"] not in used_questions]
        
        all_possible_questions.extend(unused_generic[:remaining_slots])
    
    # Final shuffle and limit to max_questions
    random.shuffle(all_possible_questions)
    return all_possible_questions[:max_questions]

def get_download_link(text, filename, label="Download"):
    b64 = base64.b64encode(text.encode()).decode()
    return f'<a href="data:file/txt;base64,{b64}" download="{filename}">{label}</a>'

def get_feedback_message(score):
    if score >= 80:
        return random.choice(EVALUATION_POSITIVE)
    elif score >= 60:
        return random.choice(EVALUATION_AVERAGE)
    else:
        return random.choice(EVALUATION_NEEDS_IMPROVEMENT)

def format_skills_message(skills):
    message = ""
    for category, skill_list in skills.items():
        message += f"**{category.capitalize()}**: {', '.join(skill_list)}\n"
    return message

def export_results_as_pdf(candidate_name, interview_date, avg_score, rating, skills, evaluations, questions):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Technical Interview Results", ln=True, align="C")
    pdf.ln(5)
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, f"Candidate: {candidate_name}", ln=True)
    pdf.cell(0, 10, f"Date: {interview_date}", ln=True)
    pdf.ln(5)
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Summary", ln=True)
    pdf.set_font("Arial", "", 12)
    pdf.cell(0, 10, f"Overall Score: {avg_score:.1f}/100", ln=True)
    pdf.cell(0, 10, f"Rating: {rating}", ln=True)
    pdf.ln(5)
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Skills", ln=True)
    pdf.set_font("Arial", "", 12)
    for category, skill_list in skills.items():
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 10, f"{category.capitalize()}", ln=True)
        pdf.set_font("Arial", "", 12)
        pdf.multi_cell(0, 10, ", ".join(skill_list))
    pdf.ln(5)
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 10, "Interview Questions and Evaluations", ln=True)
    for i, q in enumerate(questions):
        if q['question'] in evaluations:
            data = evaluations[q['question']]
            evaluation = data["evaluation"]
            pdf.set_font("Arial", "B", 12)
            pdf.cell(0, 10, f"Question {i+1}: {q['question']}", ln=True)
            pdf.set_font("Arial", "", 12)
            pdf.multi_cell(0, 10, f"Answer: {data['answer']}")
            pdf.set_font("Arial", "B", 12)
            pdf.cell(0, 10, f"Score: {evaluation.get('score', 'N/A')}/100", ln=True)
            pdf.set_font("Arial", "", 12)
            pdf.multi_cell(0, 10, f"Feedback: {evaluation.get('feedback', 'No feedback available')}")
            missing = evaluation.get('missing_concepts', [])
            if missing:
                pdf.set_font("Arial", "B", 12)
                pdf.cell(0, 10, "Missing concepts:", ln=True)
                pdf.set_font("Arial", "", 12)
                for concept in missing:
                    pdf.cell(0, 10, f"- {concept}", ln=True)
            pdf.ln(5)
    output_path = "interview_results.pdf"
    pdf.output(output_path)
    return output_path

def generate_interview_summary(candidate_name, interview_date, avg_score, rating, skills, evaluations, questions):
    summary = []
    summary.append(f"# Technical Interview Results for {candidate_name}")
    summary.append(f"**Date:** {interview_date}")
    summary.append(f"**Overall Score:** {avg_score:.1f}/100")
    summary.append(f"**Rating:** {rating}")
    summary.append("\n## Skills Profile")
    for category, skill_list in skills.items():
        summary.append(f"**{category.capitalize()}:** {', '.join(skill_list)}")
summary.append("\n## Question Analysis")
    for i, q in enumerate(questions):
        if q['question'] in evaluations:
            data = evaluations[q['question']]
            evaluation = data["evaluation"]
            score = evaluation.get('score', 0)
            summary.append(f"### Question {i+1}: {q['question']}")
            summary.append(f"**Score:** {score}/100")
            summary.append(f"**Feedback:** {evaluation.get('feedback', 'No feedback available')}")
            missing = evaluation.get('missing_concepts', [])
            if missing:
                summary.append("**Areas for improvement:**")
                for concept in missing:
                    summary.append(f"- {concept}")
            summary.append("")
    summary.append("## Interview Recommendation")
    if avg_score >= 85:
        summary.append("Based on your technical interview performance, you demonstrate strong technical knowledge and communication skills.")
    elif avg_score >= 70:
        summary.append("Your technical skills are solid, with some areas that could benefit from deeper understanding.")
    elif avg_score >= 50:
        summary.append("You have a good foundation of technical knowledge, but should continue to build your expertise.")
    else:
        summary.append("Consider spending more time studying the fundamentals of your technical areas.")
    return "\n".join(summary)

st.set_page_config(page_title="Technical Interview Chatbot", layout="wide")

if "resume_text" not in st.session_state:
    st.session_state.resume_text = ""
if "skills" not in st.session_state:
    st.session_state.skills = {}
if "questions" not in st.session_state:
    st.session_state.questions = []
if "current_question_index" not in st.session_state:
    st.session_state.current_question_index = 0
if "evaluations" not in st.session_state:
    st.session_state.evaluations = {}
if "interview_complete" not in st.session_state:
    st.session_state.interview_complete = False
if "interview_date" not in st.session_state:
    st.session_state.interview_date = datetime.now().strftime("%Y-%m-%d %H:%M")
if "chat_messages" not in st.session_state:
    st.session_state.chat_messages = [{"role": "assistant", "content": random.choice(WELCOME_MESSAGES) + " " + random.choice(RESUME_PROMPTS)}]
if "candidate_name" not in st.session_state:
    st.session_state.candidate_name = ""
if "bot_state" not in st.session_state:
    st.session_state.bot_state = "wait_for_resume"
if "max_questions" not in st.session_state:
    st.session_state.max_questions = 5
if "debug_skills" not in st.session_state:
    st.session_state.debug_skills = []
if "raw_resume_text" not in st.session_state:
    st.session_state.raw_resume_text = ""

def add_message(role, content):
    st.session_state.chat_messages.append({"role": role, "content": content})

with st.sidebar:
    st.header("Interview Bot Settings")
    if st.session_state.bot_state in ["wait_for_resume", "analyzing_resume"]:
        st.info("Please upload your resume or paste its content to begin.")
    elif st.session_state.bot_state == "interview":
        st.subheader("Interview Progress")
        progress = st.session_state.current_question_index / len(st.session_state.questions)
        st.progress(progress)
        st.write(f"Question {st.session_state.current_question_index}/{len(st.session_state.questions)}")
        if st.session_state.skills:
            st.subheader("Your Skills Focus")
            for category, skills in st.session_state.skills.items():
                with st.expander(category.capitalize()):
                    st.write(", ".join(skills))
    elif st.session_state.bot_state == "complete":
        st.success("Interview Complete!")
        evaluations = st.session_state.evaluations
        total_score = sum(data["evaluation"].get("score", 0) for data in evaluations.values())
        avg_score = total_score / len(evaluations) if evaluations else 0
        rating = "Excellent" if avg_score >= 85 else "Good" if avg_score >= 70 else "Average" if avg_score >= 50 else "Needs Improvement"
        st.metric("Overall Score", f"{avg_score:.1f}/100")
        st.metric("Rating", rating)
    
    max_q = st.slider("Number of Questions", min_value=3, max_value=10, value=st.session_state.max_questions)
    if max_q != st.session_state.max_questions:
        st.session_state.max_questions = max_q
    
    if st.button("Start New Interview"):
        st.session_state.resume_text = ""
        st.session_state.skills = {}
        st.session_state.questions = []
        st.session_state.current_question_index = 0
        st.session_state.evaluations = {}
        st.session_state.interview_complete = False
        st.session_state.bot_state = "wait_for_resume"
        st.session_state.chat_messages = [{"role": "assistant", "content": random.choice(WELCOME_MESSAGES) + " " + random.choice(RESUME_PROMPTS)}]
        st.session_state.candidate_name = ""
        st.session_state.debug_skills = []
        st.session_state.raw_resume_text = ""
        st.rerun()

st.title("Technical Interview Chatbot 🤖")

uploaded_file = st.file_uploader("Upload your resume (PDF or DOCX)", type=["pdf", "docx"], key="resume_upload")
chat_container = st.container()

def process_user_input(user_input):
    add_message("user", user_input)
    
    if st.session_state.bot_state == "wait_for_resume":
        if len(user_input.split()) <= 3 and len(st.session_state.chat_messages) <= 3:
            st.session_state.candidate_name = user_input
            add_message("assistant", f"Nice to meet you, {user_input}! Please upload your resume or paste its content so I can prepare relevant technical questions.")
            return
        
        st.session_state.resume_text = user_input
        st.session_state.bot_state = "analyzing_resume"
        add_message("assistant", "Thanks for sharing your resume! I'm analyzing it to identify your technical skills...")
        
        skills = extract_skills(user_input)
        if not skills:
            add_message("assistant", "I couldn't identify specific technical skills from your resume. Let's add some manually. What are your top technical skills? (e.g., Python, Java, AWS)")
            st.session_state.bot_state = "manual_skills"
        else:
            st.session_state.skills = skills
            skill_message = random.choice(SKILL_MESSAGES) + "\n\n" + format_skills_message(skills)
            if st.session_state.debug_skills:
                skill_message += "\n\n**Debug Info:**\n" + "\n".join(st.session_state.debug_skills)
            skill_message += f"\n\n**Raw Resume Text (first 200 chars):** {st.session_state.raw_resume_text[:200]}..."
            skill_message += "\n\nAre these skills accurate? You can add more skills if needed, or type 'start interview' when you're ready."
            add_message("assistant", skill_message)
            st.session_state.bot_state = "confirm_skills"
    
    elif st.session_state.bot_state == "manual_skills":
        skills_input = user_input.lower()
        manual_skills = {}
        for category, skill_list in COMMON_SKILLS.items():
            found_skills = []
            for skill in skill_list:
                if skill in skills_input:
                    found_skills.append(skill)
            if found_skills:
                manual_skills[category] = found_skills
        
        if not manual_skills:
            manual_skills['programming'] = ['python', 'java']
            manual_skills['tools'] = ['git']
        
        st.session_state.skills = manual_skills
        skill_message = "Thanks! I've added these skills to your profile:\n\n" + format_skills_message(manual_skills)
        skill_message += "\n\nReady to start the interview? Type 'start interview' when you're ready."
        add_message("assistant", skill_message)
        st.session_state.bot_state = "confirm_skills"
    
    elif st.session_state.bot_state == "confirm_skills":
        if "start interview" in user_input.lower() or "ready" in user_input.lower() or "yes" in user_input.lower():
            technical_questions = generate_technical_questions(st.session_state.skills, st.session_state.max_questions)
            st.session_state.questions = technical_questions
            st.session_state.current_question_index = 0
            start_message = random.choice(INTERVIEW_START_MESSAGES)
            first_question = technical_questions[0]["question"] if technical_questions else "Tell me about your background in technology."
            add_message("assistant", f"{start_message}\n\n**Question 1:** {first_question}")
            st.session_state.bot_state = "interview"
        else:
            new_skills = extract_skills(user_input)
            if new_skills:
                for category, skills_list in new_skills.items():
                    if category in st.session_state.skills:
                        for skill in skills_list:
                            if skill not in st.session_state.skills[category]:
                                st.session_state.skills[category].append(skill)
                    else:
                        st.session_state.skills[category] = skills_list
                add_message("assistant", f"I've updated your skills profile. Type 'start interview' when you're ready to begin.")
            else:
                add_message("assistant", "I'm ready whenever you are. Type 'start interview' to begin.")
    
    elif st.session_state.bot_state == "interview":
        current_question = st.session_state.questions[st.session_state.current_question_index]
        evaluation = evaluate_answer_with_nlp(current_question["question"], user_input, current_question["expected_keywords"])
        
        st.session_state.evaluations[current_question["question"]] = {
            "answer": user_input,
            "evaluation": evaluation
        }
        
        feedback = evaluation["feedback"]
        missing_concepts = evaluation.get("missing_concepts", [])
        missing_text = ""
        if missing_concepts:
            missing_text = "\n\nConsider including these concepts in your answer: " + ", ".join(missing_concepts[:3])
            if len(missing_concepts) > 3:
                missing_text += ", and others."
        
        score = evaluation.get("score", 0)
        score_message = f"\n\n**Score:** {score:.1f}/100"
        
        st.session_state.current_question_index += 1
        
        if st.session_state.current_question_index < len(st.session_state.questions):
            next_question = st.session_state.questions[st.session_state.current_question_index]
            transition = random.choice(QUESTION_TRANSITIONS)
            response = f"{feedback}{missing_text}{score_message}\n\n{transition}\n\n**Question {st.session_state.current_question_index + 1}:** {next_question['question']}"
            add_message("assistant", response)
        else:
            st.session_state.interview_complete = True
            st.session_state.bot_state = "complete"
            
            evaluations = st.session_state.evaluations
            total_score = sum(data["evaluation"].get("score", 0) for data in evaluations.values())
            avg_score = total_score / len(evaluations) if evaluations else 0
            
            if avg_score >= 85:
                rating = "Excellent"
                message = "Your technical knowledge is impressive! You demonstrated a deep understanding of the concepts."
            elif avg_score >= 70:
                rating = "Good"
                message = "You showed good technical knowledge. With a bit more practice, you'll excel in interviews."
            elif avg_score >= 50:
                rating = "Average"
                message = "You have a decent foundation, but should work on strengthening your technical knowledge."
            else:
                rating = "Needs Improvement"
                message = "You should focus on building a stronger technical foundation before your interviews."
                
            summary = generate_interview_summary(
                st.session_state.candidate_name or "Candidate", 
                st.session_state.interview_date,
                avg_score,
                rating,
                st.session_state.skills,
                evaluations,
                st.session_state.questions
            )
                
            response = f"{feedback}{missing_text}{score_message}\n\n**Interview Complete!**\n\nYour overall score is **{avg_score:.1f}/100** ({rating}).\n\n{message}\n\nHere's a summary of your performance:\n\n```\n{summary}\n```"
            add_message("assistant", response)
            
            try:
                pdf_path = export_results_as_pdf(
                    st.session_state.candidate_name or "Candidate",
                    st.session_state.interview_date,
                    avg_score,
                    rating,
                    st.session_state.skills,
                    evaluations,
                    st.session_state.questions
                )
                with open(pdf_path, "rb") as f:
                    pdf_bytes = f.read()
                pdf_b64 = base64.b64encode(pdf_bytes).decode()
                href = f'<a href="data:application/pdf;base64,{pdf_b64}" download="interview_results.pdf">Download PDF Report</a>'
                st.markdown(href, unsafe_allow_html=True)
            except Exception as e:
                st.error(f"Error generating PDF: {e}")
    
    elif st.session_state.bot_state == "complete":
        if "new interview" in user_input.lower() or "start over" in user_input.lower() or "again" in user_input.lower():
            st.session_state.resume_text = ""
            st.session_state.skills = {}
            st.session_state.questions = []
            st.session_state.current_question_index = 0
            st.session_state.evaluations = {}
            st.session_state.interview_complete = False
            st.session_state.bot_state = "wait_for_resume"
            add_message("assistant", random.choice(WELCOME_MESSAGES) + " " + random.choice(RESUME_PROMPTS))
        else:
            add_message("assistant", "The interview is complete! If you'd like to start a new interview, type 'new interview' or use the button in the sidebar.")

if uploaded_file:
    file_extension = uploaded_file.name.split(".")[-1].lower()
    
    if file_extension == "pdf":
        with st.spinner("Processing PDF..."):
            resume_text = extract_text_from_pdf(uploaded_file)
    elif file_extension == "docx":
        with st.spinner("Processing DOCX..."):
            resume_text = extract_text_from_docx(uploaded_file)
    else:
        st.error("Unsupported file format")
        resume_text = ""
    
    if resume_text and st.session_state.bot_state == "wait_for_resume":
        st.session_state.resume_text = resume_text
        st.session_state.bot_state = "analyzing_resume"
        
        add_message("user", "I've uploaded my resume.")
        add_message("assistant", "Thanks for uploading your resume! I'm analyzing it to identify your technical skills...")
        
        skills = extract_skills(resume_text)
        if not skills:
            add_message("assistant", "I couldn't identify specific technical skills from your resume. Let's add some manually. What are your top technical skills? (e.g., Python, Java, AWS)")
            st.session_state.bot_state = "manual_skills"
        else:
            st.session_state.skills = skills
            skill_message = random.choice(SKILL_MESSAGES) + "\n\n" + format_skills_message(skills)
            skill_message += "\n\nAre these skills accurate? You can add more skills if needed, or type 'start interview' when you're ready."
            add_message("assistant", skill_message)
            st.session_state.bot_state = "confirm_skills"

with chat_container:
    for message in st.session_state.chat_messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    if user_input := st.chat_input("Type here...", key="chat_input"):
        process_user_input(user_input)
        st.rerun()
